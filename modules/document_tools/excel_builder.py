"""
Excel (.xlsx) Spreadsheet Generator for Nova Smart Assistant.
Generates styled Excel workbooks with headers, borders, and auto-sized columns using openpyxl.
"""

import os
import logging
from typing import Dict, Any, List, Optional
from datetime import datetime

logger = logging.getLogger(__name__)


def get_default_output_dir() -> str:
    home = os.path.expanduser("~")
    docs = os.path.join(home, "Documents")
    if os.path.exists(docs):
        return docs
    desktop = os.path.join(home, "Desktop")
    if os.path.exists(desktop):
        return desktop
    return home


def create_excel_spreadsheet(
    sheet_title: str = "Sheet1",
    headers: Optional[List[str]] = None,
    rows: Optional[List[List[Any]]] = None,
    filename: Optional[str] = None,
    sheets: Optional[List[Dict[str, Any]]] = None,
    add_totals: bool = False,
    chart_type: Optional[str] = None
) -> Dict[str, Any]:
    """
    Create a styled Excel spreadsheet with professional formatting.
    
    Args:
        sheet_title: Name of the active worksheet (if single sheet)
        headers: List of column header names (if single sheet)
        rows: List of row data lists (if single sheet)
        filename: Optional output filename (.xlsx)
        sheets: Optional list of sheet dicts: [{"title": "...", "headers": [...], "rows": [[...], ...], "totals": bool}]
        add_totals: If True, adds a Total row with =SUM() formulas for numeric columns
        chart_type: Optional "bar" or "line" chart to embed
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")
        
        total_font = Font(name="Calibri", size=11, bold=True, color="1F497D")
        total_fill = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")

        thin_side = Side(border_style="thin", color="D9D9D9")
        double_bottom = Side(border_style="double", color="1F497D")
        cell_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)
        total_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=double_bottom)

        # Prepare sheets list
        if sheets and isinstance(sheets, list):
            sheet_configs = sheets
        else:
            sheet_configs = [{
                "title": sheet_title,
                "headers": headers or ["Column 1"],
                "rows": rows or [],
                "totals": add_totals
            }]

        total_rows_created = 0
        for s_idx, s_cfg in enumerate(sheet_configs):
            s_title = (s_cfg.get("title") or f"Sheet{s_idx+1}")[:30]
            s_headers = s_cfg.get("headers", ["Column 1"])
            s_rows = s_cfg.get("rows", [])
            s_totals = s_cfg.get("totals", add_totals)

            if s_idx == 0:
                ws = wb.active
                ws.title = s_title
            else:
                ws = wb.create_sheet(title=s_title)

            # Freeze top row
            ws.freeze_panes = "A2"

            # Write headers
            ws.append(s_headers)
            for col_num in range(1, len(s_headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = cell_border

            # Write data rows
            for r_data in s_rows:
                ws.append(r_data)
                total_rows_created += 1

            num_rows = len(s_rows)
            # Apply formatting
            for row_idx in range(2, num_rows + 2):
                for col_idx in range(1, len(s_headers) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = cell_border
                    if isinstance(cell.value, (int, float)):
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.alignment = Alignment(horizontal="left")

            # Apply Auto-Filter
            if s_headers:
                last_col_letter = get_column_letter(len(s_headers))
                ws.auto_filter.ref = f"A1:{last_col_letter}{num_rows + 1}"

            # Calculate Totals row if requested
            if s_totals and num_rows > 0:
                total_row_idx = num_rows + 2
                ws.cell(row=total_row_idx, column=1, value="Total").font = total_font
                ws.cell(row=total_row_idx, column=1).alignment = Alignment(horizontal="left")
                ws.cell(row=total_row_idx, column=1).fill = total_fill
                ws.cell(row=total_row_idx, column=1).border = total_border

                for col_idx in range(2, len(s_headers) + 1):
                    col_letter = get_column_letter(col_idx)
                    # Check if column has numbers
                    first_val = ws.cell(row=2, column=col_idx).value
                    if isinstance(first_val, (int, float)):
                        t_cell = ws.cell(
                            row=total_row_idx, 
                            column=col_idx, 
                            value=f"=SUM({col_letter}2:{col_letter}{num_rows + 1})"
                        )
                        t_cell.font = total_font
                        t_cell.fill = total_fill
                        t_cell.border = total_border
                        t_cell.alignment = Alignment(horizontal="right")
                    else:
                        t_cell = ws.cell(row=total_row_idx, column=col_idx, value="")
                        t_cell.fill = total_fill
                        t_cell.border = total_border

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value is not None:
                        val_str = str(cell.value)
                        if len(val_str) > max_len:
                            max_len = len(val_str)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

            # Optional Chart
            if chart_type and chart_type.lower() in ("bar", "line", "column") and num_rows > 1 and len(s_headers) >= 2:
                try:
                    from openpyxl.chart import BarChart, LineChart, Reference
                    if chart_type.lower() == "line":
                        chart = LineChart()
                        chart.title = f"{s_title} Trend"
                    else:
                        chart = BarChart()
                        chart.type = "col"
                        chart.style = 10
                        chart.title = f"{s_title} Overview"

                    data_ref = Reference(ws, min_col=2, min_row=1, max_col=len(s_headers), max_row=num_rows + 1)
                    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=num_rows + 1)
                    chart.add_data(data_ref, titles_from_data=True)
                    chart.set_categories(cats_ref)
                    chart.width = 16
                    chart.height = 10
                    ws.add_chart(chart, f"{get_column_letter(len(s_headers) + 2)}2")
                except Exception as c_err:
                    logger.warning(f"[EXCEL] Could not generate chart: {c_err}")

        # Determine target file path
        first_title = sheet_configs[0]["title"]
        if not filename:
            safe_title = "".join(c for c in first_title if c.isalnum() or c in (' ', '_', '-')).rstrip()
            filename = f"{safe_title.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        elif not filename.endswith(".xlsx"):
            filename += ".xlsx"

        if os.path.isabs(filename):
            target_path = filename
            out_dir = os.path.dirname(target_path)
        else:
            out_dir = get_default_output_dir()
            target_path = os.path.join(out_dir, filename)

        os.makedirs(os.path.dirname(os.path.abspath(target_path)), exist_ok=True)
        wb.save(target_path)
        logger.info(f"[EXCEL] Spreadsheet saved to: {target_path}")

        return {
            "success": True,
            "status": "success",
            "filepath": target_path,
            "filename": os.path.basename(target_path),
            "sheets_count": len(sheet_configs),
            "total_rows": total_rows_created,
            "message": f"Excel workbook '{os.path.basename(target_path)}' with {len(sheet_configs)} sheet(s) and {total_rows_created} rows created in {out_dir}."
        }
    except Exception as e:
        logger.error(f"[EXCEL] Failed to create spreadsheet: {e}")
        return {
            "success": False,
            "status": "error",
            "error": str(e),
            "message": f"Failed to create Excel spreadsheet: {e}"
        }
